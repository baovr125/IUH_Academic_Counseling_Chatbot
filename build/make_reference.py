"""Dựng build/reference.docx theo quy định định dạng của Khoa CNTT – IUH.

Nguồn quy định: "Mau bao cao KLTN.docx.pdf" (trang 15–19) và file mẫu
04.-MAU-KHOA-LUAN-TOT-NGHIEP-2021.docx (lề, tên style). Script xuất
reference.docx mặc định của pandoc rồi ghi đè trang giấy và style, nên chạy
lại bao nhiêu lần cũng cho cùng một kết quả.

Chạy:  python build/make_reference.py
"""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import docx
import pypandoc
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm

BUILD = Path(__file__).resolve().parent
OUT = BUILD / "reference.docx"

FONT = "Times New Roman"
CODE_FONT = "Consolas"
TEXT_WIDTH_CM = 21.0 - 3.5 - 2.0  # 15,5 cm vùng nội dung


def tw(cm: float) -> int:
    """cm -> twip (1/20 pt)."""
    return round(cm / 2.54 * 1440)


# ---------------------------------------------------------------------------
# Đặc tả style. Đơn vị: cỡ chữ pt, khoảng cách pt, giãn dòng là bội số dòng.
# ---------------------------------------------------------------------------
BASE = dict(font=FONT, size=13, color="000000")

STYLES: list[dict] = [
    # Normal: dùng cho các danh mục (mục lục, danh mục hình/bảng): 0/0, 1,15
    dict(id="Normal", name="Normal", before=0, after=0, line=1.15, **BASE),
    # Nội dung (PDF tr.17): TNR 13, trước 10 / sau 0, 1,5, căn đều, không thụt
    dict(id="Content", name="Content", base="Normal", next="Content",
         before=10, after=0, line=1.5, jc="both"),
    dict(id="BodyText", name="Body Text", base="Content", next="Content"),
    dict(id="FirstParagraph", name="First Paragraph", base="Content", next="Content"),
    dict(id="Compact", name="Compact", base="Normal", before=0, after=0, line=1.15),
    # Bullet (PDF tr.19): 0/0, 1,5, căn trái
    dict(id="Bullet", name="Bullet", base="Normal", before=0, after=0, line=1.5, jc="left"),
    # Heading 1 (PDF tr.17): 14, đậm, IN HOA, 24/24, 1,15, căn trái, sang trang mới
    dict(id="Heading1", name="heading 1", base="Normal", next="Content",
         size=14, bold=True, caps=True, before=24, after=24, line=1.15, jc="left",
         keep_next=True, keep_lines=True, page_break_before=True, outline=0,
         tabs=[("left", 3.0)]),
    # Heading 1 không đánh số nhưng vào mục lục (TÀI LIỆU THAM KHẢO, PHỤ LỤC…)
    dict(id="Heading1N", name="Heading 1N", base="Heading1", next="Content", jc="center",
         outline=0),
    # Tiêu đề phần đầu (LỜI CẢM ƠN, MỤC LỤC…): không đánh số, không vào mục lục
    dict(id="Heading1Front", name="Heading 1 Front", base="Heading1", next="Content",
         jc="center", outline=9),
    # Heading 2 (PDF tr.17): 13, đậm, 6/12, 1,15
    dict(id="Heading2", name="heading 2", base="Normal", next="Content",
         size=13, bold=True, before=6, after=12, line=1.15, jc="left",
         keep_next=True, keep_lines=True, outline=1, tabs=[("left", 1.0)]),
    dict(id="Heading2N", name="Heading 2N", base="Heading2", next="Content", outline=1),
    # Heading 3: đậm nghiêng
    dict(id="Heading3", name="heading 3", base="Normal", next="Content",
         size=13, bold=True, italic=True, before=6, after=12, line=1.15, jc="left",
         keep_next=True, keep_lines=True, outline=2, tabs=[("left", 1.25)]),
    # Heading 4: nghiêng, không đậm; không vào mục lục (TOC chỉ lấy cấp 1–3)
    dict(id="Heading4", name="heading 4", base="Normal", next="Content",
         size=13, bold=False, italic=True, before=6, after=12, line=1.15, jc="left",
         keep_next=True, keep_lines=True, outline=3, tabs=[("left", 1.5)]),
    # Caption (PDF tr.18): 13, 6/12, 1,15, căn giữa
    dict(id="Caption", name="caption", base="Normal", next="Content",
         size=13, italic=False, bold=False, before=6, after=12, line=1.15, jc="center"),
    dict(id="ImageCaption", name="Image Caption", base="Caption"),
    dict(id="TableCaption", name="Table Caption", base="Caption", keep_next=True),
    # Đoạn chứa hình: căn giữa, dính với caption bên dưới
    dict(id="Figure", name="Figure", base="Normal", before=6, after=0, line=1.0,
         jc="center", keep_next=True),
    dict(id="CaptionedFigure", name="Captioned Figure", base="Figure"),
    # Công thức trong bảng 3 cột (giữa: công thức, phải: số hiệu)
    dict(id="Equation", name="Equation", base="Normal", before=6, after=6, line=1.0, jc="center"),
    dict(id="EquationNumber", name="Equation Number", base="Normal", before=6, after=6,
         line=1.0, jc="right"),
    # Chữ trong bảng (PDF tr.20: 10–13 pt, thống nhất)
    dict(id="TableText", name="Table Text", base="Normal", size=12, before=3, after=3,
         line=1.15, jc="left"),
    # Mã nguồn: Consolas 10, khung viền, nền xám nhạt
    dict(id="SourceCode", name="Source Code", base="Normal", font=CODE_FONT, size=10,
         before=6, after=6, line=1.0, jc="left", box=True, shade="F2F2F2"),
    # Ghi chú cần bổ sung ([TODO], [CHỜ SỐ LIỆU]) — tô nền vàng để dễ thấy
    dict(id="TODO", name="TODO", base="Content", italic=True, shade="FFF2CC", box=True),
    dict(id="BlockText", name="Block Text", base="Content", italic=True, ind_left=1.0),
    # Hộp "Ví dụ" (kế hoạch mục 6.1): các đoạn liền nhau cùng style được Word gộp thành một khung
    dict(id="ExampleBox", name="Example Box", base="Content", before=4, after=4, line=1.3,
         box=True, shade="F1F5F9"),
    # Tài liệu tham khảo: thụt treo 1 cm
    dict(id="Reference", name="Reference", base="Content", before=6, ind_left=1.0,
         hanging=1.0),
    # Danh mục: TNR 13, 0/0, 1,15 (PDF tr.19)
    dict(id="TOCHeading", name="TOC Heading", base="Heading1Front"),
    dict(id="TOC1", name="toc 1", base="Normal", next="Normal", before=6, after=0, line=1.15,
         tabs=[("right", TEXT_WIDTH_CM, "dot")]),
    dict(id="TOC2", name="toc 2", base="Normal", next="Normal", before=0, after=0, line=1.15,
         ind_left=0.5, tabs=[("right", TEXT_WIDTH_CM, "dot")]),
    dict(id="TOC3", name="toc 3", base="Normal", next="Normal", before=0, after=0, line=1.15,
         ind_left=1.0, tabs=[("right", TEXT_WIDTH_CM, "dot")]),
    dict(id="TableofFigures", name="table of figures", base="Normal", next="Normal",
         before=0, after=0, line=1.15, tabs=[("right", TEXT_WIDTH_CM, "dot")]),
    # Header / footer
    # Chỉ một tab phải: tên sinh viên bên trái, số trang sát lề phải
    dict(id="Header", name="header", base="Normal", size=12, line=1.0,
         tabs=[("right", TEXT_WIDTH_CM)]),
    dict(id="Footer", name="footer", base="Normal", size=12, line=1.0,
         tabs=[("right", TEXT_WIDTH_CM)]),
    dict(id="FootnoteText", name="footnote text", base="Normal", size=11, line=1.0),
    # Các style pandoc có sẵn nhưng không dùng: đưa về TNR cho an toàn
    dict(id="Title", name="Title", base="Normal", size=16, bold=True, jc="center"),
    dict(id="Subtitle", name="Subtitle", base="Normal", size=14, jc="center"),
    # Bìa và phụ bìa (mẫu của Khoa: TNR 13, căn giữa; tên đề tài 16 đậm)
    dict(id="CoverText", name="Cover Text", base="Normal", size=13, before=0, after=6, line=1.15, jc="center"),
    dict(id="CoverBold", name="Cover Bold", base="CoverText", bold=True),
    dict(id="CoverTitle", name="Cover Title", base="CoverText", size=16, bold=True, line=1.25),
    dict(id="Abstract", name="Abstract", base="Content"),
    dict(id="Bibliography", name="Bibliography", base="Reference"),
]

CHAR_STYLES: list[dict] = [
    dict(id="VerbatimChar", name="Verbatim Char", font=CODE_FONT, size=11),
]


# ---------------------------------------------------------------------------
def rpr_xml(s: dict) -> str:
    parts = []
    if "font" in s:
        f = s["font"]
        parts.append(f'<w:rFonts w:ascii="{f}" w:hAnsi="{f}" w:eastAsia="{f}" w:cs="{f}"/>')
    for key, tag in (("bold", "b"), ("italic", "i")):
        if key in s:
            v = "" if s[key] else ' w:val="0"'
            parts.append(f"<w:{tag}{v}/><w:{tag}Cs{v}/>")
    if "caps" in s:
        parts.append("<w:caps/>" if s["caps"] else '<w:caps w:val="0"/>')
    if "color" in s:
        parts.append(f'<w:color w:val="{s["color"]}"/>')
    if "size" in s:
        hp = round(s["size"] * 2)
        parts.append(f'<w:sz w:val="{hp}"/><w:szCs w:val="{hp}"/>')
    if not parts:
        return ""
    return f"<w:rPr {nsdecls('w')}>{''.join(parts)}</w:rPr>"


def ppr_xml(s: dict) -> str:
    p = []
    if s.get("keep_next"):
        p.append("<w:keepNext/>")
    if s.get("keep_lines"):
        p.append("<w:keepLines/>")
    if "page_break_before" in s:
        p.append("<w:pageBreakBefore/>" if s["page_break_before"] else '<w:pageBreakBefore w:val="0"/>')
    if s.get("box"):
        b = '<w:{0} w:val="single" w:sz="4" w:space="4" w:color="808080"/>'
        p.append("<w:pBdr>" + "".join(b.format(x) for x in ("top", "left", "bottom", "right")) + "</w:pBdr>")
    if s.get("shade"):
        p.append(f'<w:shd w:val="clear" w:color="auto" w:fill="{s["shade"]}"/>')
    if s.get("tabs"):
        t = []
        for tab in s["tabs"]:
            kind, pos = tab[0], tab[1]
            leader = f' w:leader="{tab[2]}"' if len(tab) > 2 else ""
            t.append(f'<w:tab w:val="{kind}"{leader} w:pos="{tw(pos)}"/>')
        p.append("<w:tabs>" + "".join(t) + "</w:tabs>")
    sp = []
    if "before" in s:
        sp.append(f'w:before="{round(s["before"] * 20)}"')
    if "after" in s:
        sp.append(f'w:after="{round(s["after"] * 20)}"')
    if "line" in s:
        sp.append(f'w:line="{round(s["line"] * 240)}" w:lineRule="auto"')
    if sp:
        p.append(f"<w:spacing {' '.join(sp)}/>")
    if any(k in s for k in ("ind_left", "hanging")):
        attrs = [f'w:left="{tw(s.get("ind_left", 0))}"']
        attrs.append(f'w:hanging="{tw(s["hanging"])}"' if "hanging" in s else 'w:firstLine="0"')
        p.append(f"<w:ind {' '.join(attrs)}/>")
    if "jc" in s:
        p.append(f'<w:jc w:val="{s["jc"]}"/>')
    if "outline" in s:
        p.append(f'<w:outlineLvl w:val="{s["outline"]}"/>')
    if not p:
        return ""
    return f"<w:pPr {nsdecls('w')}>{''.join(p)}</w:pPr>"


# Style dựng sẵn của Word: không được gắn customStyle, nếu không Word sẽ bỏ qua định
# dạng của ta và dùng định nghĩa mặc định (ví dụ tab của footer, style mục lục).
BUILTIN_NAMES = {"normal", "caption", "header", "footer", "footnote text", "title", "subtitle",
                 "block text", "body text", "table of figures", "toc heading",
                 *(f"heading {i}" for i in range(1, 10)), *(f"toc {i}" for i in range(1, 10))}


def ensure_style(styles_el, sid: str, name: str, stype: str = "paragraph"):
    for st in styles_el.findall(qn("w:style")):
        if st.get(qn("w:styleId")) == sid:
            break
    else:
        st = parse_xml(f'<w:style {nsdecls("w")} w:type="{stype}" w:styleId="{sid}">'
                       f'<w:name w:val="{name}"/><w:qFormat/></w:style>')
        styles_el.append(st)
    if name.lower() in BUILTIN_NAMES:
        st.attrib.pop(qn("w:customStyle"), None)
    else:
        st.set(qn("w:customStyle"), "1")
    return st


def apply_style(styles_el, s: dict, stype: str = "paragraph"):
    st = ensure_style(styles_el, s["id"], s["name"], stype)
    # Xóa định dạng cũ (pandoc dùng phông theme, màu xanh cho heading)
    for tag in ("w:basedOn", "w:next", "w:link", "w:pPr", "w:rPr"):
        for el in st.findall(qn(tag)):
            st.remove(el)
    name_el = st.find(qn("w:name"))
    name_el.set(qn("w:val"), s["name"])
    pos = list(st).index(name_el) + 1
    if s.get("base"):
        st.insert(pos, parse_xml(f'<w:basedOn {nsdecls("w")} w:val="{s["base"]}"/>'))
        pos += 1
    if s.get("next"):
        st.insert(pos, parse_xml(f'<w:next {nsdecls("w")} w:val="{s["next"]}"/>'))
    if stype == "paragraph":
        x = ppr_xml(s)
        if x:
            st.append(parse_xml(x))
    x = rpr_xml(s)
    if x:
        st.append(parse_xml(x))


def strip_theme_fonts(root):
    """Bỏ các thuộc tính *Theme của rFonts để phông TNR không bị theme ghi đè."""
    for rf in root.iter(qn("w:rFonts")):
        for attr in list(rf.attrib):
            if attr.endswith("Theme}") or "Theme" in attr:
                del rf.attrib[attr]


def set_doc_defaults(styles_el):
    dd = styles_el.find(qn("w:docDefaults"))
    if dd is not None:
        styles_el.remove(dd)
    dd = parse_xml(
        f'<w:docDefaults {nsdecls("w")}>'
        f'<w:rPrDefault><w:rPr><w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" w:eastAsia="{FONT}" w:cs="{FONT}"/>'
        '<w:sz w:val="26"/><w:szCs w:val="26"/><w:lang w:val="vi-VN" w:eastAsia="en-US" w:bidi="ar-SA"/>'
        '</w:rPr></w:rPrDefault>'
        '<w:pPrDefault><w:pPr><w:spacing w:after="0" w:line="276" w:lineRule="auto"/></w:pPr></w:pPrDefault>'
        '</w:docDefaults>')
    styles_el.insert(0, dd)


def main() -> int:
    tmp = BUILD / "_pandoc_default_reference.docx"
    pandoc = pypandoc.get_pandoc_path()
    subprocess.run([pandoc, "-o", str(tmp), "--print-default-data-file", "reference.docx"], check=True)

    d = docx.Document(str(tmp))
    # Trang giấy: A4, lề trên 3, dưới 3, trái 3,5, phải 2 (PDF tr.15)
    for sec in d.sections:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.top_margin, sec.bottom_margin = Cm(3.0), Cm(3.0)
        sec.left_margin, sec.right_margin = Cm(3.5), Cm(2.0)
        sec.header_distance = sec.footer_distance = Cm(1.27)

    styles_el = d.styles.element
    set_doc_defaults(styles_el)
    for s in STYLES:
        apply_style(styles_el, s)
    for s in CHAR_STYLES:
        apply_style(styles_el, s, "character")
    strip_theme_fonts(styles_el)
    # Heading pandoc mặc định có màu xanh ở style ký tự liên kết (Heading 1 Char…)
    for st in styles_el.findall(qn("w:style")):
        for color in st.iter(qn("w:color")):
            if st.get(qn("w:type")) == "character" and st.get(qn("w:styleId")).startswith("Heading"):
                color.set(qn("w:val"), "000000")

    # Nội dung mẫu để mở reference.docx ra là thấy ngay style
    body = d.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    for sid, text in (("Heading1", "Tên chương (Heading 1)"), ("Heading2", "Tiểu mục (Heading 2)"),
                      ("Heading3", "Tiểu mục (Heading 3)"), ("Heading4", "Tiểu mục (Heading 4)"),
                      ("Content", "Đoạn nội dung (Content): Times New Roman 13, giãn dòng 1,5."),
                      ("Bullet", "Gạch đầu dòng (Bullet)"), ("Caption", "Hình 1.1: Chú thích (Caption)")):
        p = d.add_paragraph(text)
        p._p.get_or_add_pPr().insert(0, parse_xml(f'<w:pStyle {nsdecls("w")} w:val="{sid}"/>'))
        # add_paragraph đặt đoạn trước sectPr nên thứ tự được giữ
    d.save(str(OUT))
    tmp.unlink(missing_ok=True)
    print(f"Đã tạo {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
