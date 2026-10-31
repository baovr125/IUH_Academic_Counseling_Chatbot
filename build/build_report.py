"""Ghép bản thảo Markdown (Report/drafts) thành DOCX theo định dạng của Khoa CNTT – IUH.

Quy trình:
  1. Đọc build/report_config.json, ghép các file .md theo 3 phần
     front (phần đầu) → main (Chương 1 → Tài liệu tham khảo) → appendix (phụ lục).
  2. pandoc + build/reference.docx + build/filters.lua → DOCX thô.
  3. Hậu xử lý bằng python-docx: đánh số heading "CHƯƠNG x / x.y / x.y.z / x.y.z.t",
     đổi style đoạn (Content, Bullet, Table Text), bảng chỉ kẻ ngang, công thức có số
     căn phải, thu nhỏ hình quá khổ, chia section + header/footer + số trang,
     bật cập nhật trường khi mở file.
  4. (--word) Dùng Word cập nhật mục lục, danh mục, số caption và xuất PDF.

Chạy (từ thư mục Report):
  python build/build_report.py            # báo cáo thật từ drafts/
  python build/build_report.py --sample   # file mẫu kiểm tra định dạng
  python build/build_report.py --sample --word --pdf
"""
from __future__ import annotations

import argparse
import copy
import json
import os
import re
import subprocess
import sys
from pathlib import Path

import docx
import pypandoc
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm

BUILD = Path(__file__).resolve().parent
REPORT = BUILD.parent
NS_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
NS_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
MARK_RE = re.compile(r"@@SECTION_END:(\w+)@@")
TEXT_WIDTH = Cm(15.5)
MAX_FIGURE_HEIGHT = Cm(20)

warnings: list[str] = []


def warn(msg: str) -> None:
    warnings.append(msg)
    print(f"[CẢNH BÁO] {msg}", file=sys.stderr)


# ---------------------------------------------------------------------------
# Tiện ích XML
# ---------------------------------------------------------------------------
def pstyle(p) -> str | None:
    ps = p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
    return None if ps is None else ps.get(qn("w:val"))


def get_ppr(p):
    ppr = p.find(qn("w:pPr"))
    if ppr is None:
        ppr = parse_xml(f"<w:pPr {nsdecls('w')}/>")
        p.insert(0, ppr)
    return ppr


def set_pstyle(p, sid: str) -> None:
    ppr = get_ppr(p)
    ps = ppr.find(qn("w:pStyle"))
    if ps is None:
        ps = parse_xml(f'<w:pStyle {nsdecls("w")} w:val="{sid}"/>')
        ppr.insert(0, ps)
    ps.set(qn("w:val"), sid)


def ptext(p) -> str:
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


def insert_before_first(parent, new_el, tags: tuple[str, ...]) -> None:
    """Chèn new_el trước phần tử con đầu tiên có tag thuộc `tags` (giữ đúng thứ tự schema)."""
    wanted = {qn("w:" + t) for t in tags}
    for child in parent:
        if child.tag in wanted:
            child.addprevious(new_el)
            return
    parent.append(new_el)


def has_ancestor(el, tag: str) -> bool:
    a = el.getparent()
    while a is not None:
        if a.tag == qn(tag):
            return True
        a = a.getparent()
    return False


# ---------------------------------------------------------------------------
# 1. Ghép Markdown
# ---------------------------------------------------------------------------
IMG_RE = re.compile(r"(!\[(?:[^\[\]]|\[[^\]]*\])*\]\()(<[^>]+>|[^)\s]+)")
FIG_ID_RE = re.compile(r"^H\d-\d\d[a-z]?$")
_catalog: dict[str, str] | None = None


def catalog_captions() -> dict[str, str]:
    """Chú thích chuẩn trong figures/FIGURES.md, theo mã hình."""
    global _catalog
    if _catalog is None:
        _catalog = {}
        path = REPORT / "figures" / "FIGURES.md"
        if path.exists():
            for line in path.read_text(encoding="utf-8").splitlines():
                cells = [c.strip() for c in line.strip().strip("|").split("|")]
                if len(cells) == 7 and FIG_ID_RE.match(cells[0]):
                    _catalog[cells[0]] = cells[4]
    return _catalog


def _norm_caption(text: str) -> str:
    text = re.sub(r"^Hình\s+[\dxA-Z]+\.[\dxy]+[.:]?\s*", "", text.strip())
    return re.sub(r"\s+", " ", text.replace("`", "")).strip().rstrip(".")


def preprocess(md_path: Path) -> str:
    """Đổi đường dẫn ảnh tương đối (tính từ file .md) thành tương đối với thư mục Report,
    và cảnh báo khi chú thích hình khác chú thích chuẩn trong FIGURES.md."""
    text = md_path.read_text(encoding="utf-8")

    def fix(m: re.Match) -> str:
        url = m.group(2).strip("<>")
        if re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*://", url):
            return m.group(0)
        target = (md_path.parent / url).resolve()
        if not target.exists():
            warn(f"{md_path.name}: không tìm thấy ảnh {url}")
        fid = target.stem
        if FIG_ID_RE.match(fid):
            alt = m.group(1)[2:-2]
            expected = catalog_captions().get(fid)
            if expected is None:
                warn(f"{md_path.name}: hình {fid} không có trong FIGURES.md")
            elif _norm_caption(alt) != _norm_caption(expected):
                warn(f"{md_path.name}: chú thích {fid} khác FIGURES.md\n    bản thảo : {_norm_caption(alt)}"
                     f"\n    FIGURES  : {_norm_caption(expected)}")
        rel = os.path.relpath(target, REPORT).replace("\\", "/")
        return f"{m.group(1)}<{rel}>"

    return IMG_RE.sub(fix, text)


def fill_placeholders(markdown: str, cfg: dict) -> str:
    """Thay {{KHÓA}} trong bản thảo bìa bằng thông tin ở report_config.json (mục "cover" và "students")."""
    values = {k: str(v) for k, v in cfg.get("cover", {}).items()}
    for i, name in enumerate(cfg.get("students", []), 1):
        values[f"sv{i}"] = name
        values[f"SV{i}"] = name.upper()
    return re.sub(r"\{\{(\w+)\}\}", lambda m: values.get(m.group(1), m.group(0)), markdown)


def assemble(parts: dict[str, list[str]]) -> tuple[str, list[str]]:
    chunks, names = [], []
    present = []
    for name in ("cover", "front", "main", "appendix"):
        files = []
        for f in parts.get(name, []):
            path = REPORT / f
            if path.exists():
                files.append(path)
            else:
                print(f"[bỏ qua] chưa có {f}")
        if files:
            present.append((name, files))
    for idx, (name, files) in enumerate(present):
        for path in files:
            chunks.append(preprocess(path))
        if idx < len(present) - 1:
            chunks.append(f"```{{=openxml}}\n<w:p><w:r><w:t>@@SECTION_END:{name}@@</w:t></w:r></w:p>\n```")
        names.append(name)
    return resolve_citations("\n\n".join(chunks) + "\n"), names


# ---------------------------------------------------------------------------
# Trích dẫn theo khóa
# ---------------------------------------------------------------------------
# Trong danh mục tài liệu tham khảo, mỗi mục mở đầu bằng khóa:  [@gao2023] Gao Y. và cs. (2023), ...
# Trong văn bản trích dẫn bằng [@gao2023] hoặc [@gao2023; @lewis2020]. Khi ghép, khóa được đổi thành
# số thứ tự [n] theo đúng thứ tự xuất hiện trong danh mục (danh mục do nhóm sắp xếp theo quy định 2.3).
REF_DEF_RE = re.compile(r"^\[@([\w:.-]+)\][ \t]+", re.MULTILINE)
CITE_RE = re.compile(r"\[(@[\w:.-]+(?:\s*;\s*@[\w:.-]+)*)\]")


def _format_numbers(nums: list[int]) -> str:
    """[1], [3]–[5]: gộp dãy liên tiếp từ 3 số trở lên thành khoảng."""
    nums = sorted(set(nums))
    parts, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        if j - i >= 2:
            parts.append(f"[{nums[i]}]–[{nums[j]}]")
        else:
            parts.extend(f"[{n}]" for n in nums[i:j + 1])
        i = j + 1
    return ", ".join(parts)


def resolve_citations(text: str) -> str:
    keys = REF_DEF_RE.findall(text)
    if not keys:
        return text
    dup = {k for k in keys if keys.count(k) > 1}
    if dup:
        warn(f"Khóa tài liệu tham khảo bị trùng: {', '.join(sorted(dup))}")
    number = {}
    for k in keys:
        number.setdefault(k, len(number) + 1)
    used: set[str] = set()

    def cite(m: re.Match) -> str:
        found = []
        for key in re.findall(r"@([\w:.-]+)", m.group(1)):
            if key in number:
                found.append(number[key])
                used.add(key)
            else:
                warn(f"Trích dẫn [@{key}] không có trong danh mục tài liệu tham khảo")
        return _format_numbers(found) if found else m.group(0)

    text = REF_DEF_RE.sub(lambda m: f"[{number[m.group(1)]}]. ", text)
    text = CITE_RE.sub(cite, text)
    unused = [k for k in number if k not in used]
    if unused:
        print(f"[ghi chú] {len(unused)} tài liệu chưa được trích dẫn: {', '.join(unused)}")
    print(f"[trích dẫn] {len(number)} tài liệu, {len(used)} được trích dẫn")
    return text


def run_pandoc(markdown: str, out: Path, cfg: dict) -> None:
    pandoc = pypandoc.get_pandoc_path()
    # lists_without_preceding_blankline: cho phép danh sách đứng ngay sau dòng "trong đó:" mà không cần dòng trống
    args = [pandoc, "-f", "markdown+lists_without_preceding_blankline", "-t", "docx", "-o", str(out),
            f"--reference-doc={BUILD / 'reference.docx'}",
            f"--lua-filter={BUILD / 'filters.lua'}",
            "--syntax-highlighting=none",
            "-M", f"caption-sep={cfg.get('caption_sep', 'colon')}"]
    res = subprocess.run(args, input=markdown.encode("utf-8"), cwd=REPORT,
                         capture_output=True)
    err = res.stderr.decode("utf-8", errors="replace").strip()
    if err:
        print(err, file=sys.stderr)
        warnings.extend(l for l in err.splitlines() if l.startswith("[CẢNH BÁO]"))
    if res.returncode != 0:
        raise SystemExit(f"pandoc lỗi (mã {res.returncode})")


# ---------------------------------------------------------------------------
# 2. Hậu xử lý
# ---------------------------------------------------------------------------
# (mẫu số, style, vị trí tab sau số – cm). Dòng thứ hai của tiêu đề dài thụt treo đúng
# bằng vị trí tab để thẳng hàng với chữ đầu tiên của tiêu đề.
HEADING_LEVELS = [("CHƯƠNG %1", "Heading1", 3.0), ("%1.%2", "Heading2", 1.0),
                  ("%1.%2.%3", "Heading3", 1.25), ("%1.%2.%3.%4", "Heading4", 1.5)]
PPR_AFTER_NUMPR = ("suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
                   "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
                   "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
                   "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
                   "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl",
                   "divId", "cnfStyle", "rPr", "sectPr", "pPrChange")


def style_el(doc, sid: str):
    for st in doc.styles.element.findall(qn("w:style")):
        if st.get(qn("w:styleId")) == sid:
            return st
    return None


def set_style_numpr(doc, sid: str, num_id: int, ilvl: int | None) -> None:
    st = style_el(doc, sid)
    if st is None:
        return
    ppr = st.find(qn("w:pPr"))
    if ppr is None:
        ppr = parse_xml(f"<w:pPr {nsdecls('w')}/>")
        rpr = st.find(qn("w:rPr"))
        (rpr.addprevious(ppr) if rpr is not None else st.append(ppr))
    for old in ppr.findall(qn("w:numPr")):
        ppr.remove(old)
    lvl = "" if ilvl is None else f'<w:ilvl w:val="{ilvl}"/>'
    insert_before_first(ppr, parse_xml(f'<w:numPr {nsdecls("w")}>{lvl}<w:numId w:val="{num_id}"/></w:numPr>'),
                        PPR_AFTER_NUMPR)


def add_heading_numbering(doc, layout: str) -> None:
    numbering = doc.part.numbering_part.element
    abs_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    aid, nid = max(abs_ids, default=0) + 1, max(num_ids, default=0) + 1
    lvls = []
    for i, (text, sid, tab_cm) in enumerate(HEADING_LEVELS):
        two_lines = i == 0 and layout == "two_lines_center"
        suff = "nothing" if two_lines else "tab"
        ind = '<w:ind w:left="0" w:firstLine="0"/>' if two_lines else \
            f'<w:ind w:left="{round(tab_cm / 2.54 * 1440)}" w:hanging="{round(tab_cm / 2.54 * 1440)}"/>'
        lgl = "<w:isLgl/>" if i else ""
        lvls.append(f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
                    f'<w:pStyle w:val="{sid}"/>{lgl}<w:suff w:val="{suff}"/>'
                    f'<w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>'
                    f'<w:pPr>{ind}</w:pPr></w:lvl>')
    absnum = parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{aid}">'
                       f'<w:multiLevelType w:val="multilevel"/>{"".join(lvls)}</w:abstractNum>')
    first_num = numbering.find(qn("w:num"))
    (first_num.addprevious(absnum) if first_num is not None else numbering.append(absnum))
    num = parse_xml(f'<w:num {nsdecls("w")} w:numId="{nid}"><w:abstractNumId w:val="{aid}"/></w:num>')
    cleanup = numbering.find(qn("w:numIdMacAtCleanup"))
    (cleanup.addprevious(num) if cleanup is not None else numbering.append(num))

    for i, (_, sid, _) in enumerate(HEADING_LEVELS):
        set_style_numpr(doc, sid, nid, i)
    for sid in ("Heading1N", "Heading1Front", "Heading2N", "TOCHeading"):
        set_style_numpr(doc, sid, 0, None)

    if layout == "two_lines_center":
        st = style_el(doc, "Heading1")
        st.find(f"{qn('w:pPr')}/{qn('w:jc')}").set(qn("w:val"), "center")
        for p in doc.element.body.iter(qn("w:p")):
            if pstyle(p) == "Heading1":
                first_run = p.find(qn("w:r"))
                br = parse_xml(f"<w:r {nsdecls('w')}><w:br/></w:r>")
                (first_run.addprevious(br) if first_run is not None else p.append(br))


def restyle_paragraphs(body) -> None:
    for p in body.iter(qn("w:p")):
        st = pstyle(p)
        if p.find(f".//{{{NS_M}}}oMathPara") is not None:
            set_pstyle(p, "Equation")
        elif has_ancestor(p, "w:tbl"):
            if st in (None, "Normal", "Compact", "BodyText", "FirstParagraph"):
                set_pstyle(p, "TableText")
        elif p.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is not None and st in ("Compact", "BodyText", "FirstParagraph", None):
            set_pstyle(p, "Bullet")
        elif st in ("BodyText", "FirstParagraph", "Compact"):
            set_pstyle(p, "Content")


def format_tables(body) -> int:
    """Bảng chỉ kẻ đường ngang, rộng bằng vùng nội dung, dòng tiêu đề in đậm và lặp lại."""
    n = 0
    for tbl in list(body.iter(qn("w:tbl"))):
        n += 1
        tblpr = tbl.find(qn("w:tblPr"))
        keep = {t: tblpr.find(qn(t)) for t in ("w:tblStyle", "w:tblCellMar", "w:tblLook")}
        for child in list(tblpr):
            tblpr.remove(child)
        if keep["w:tblStyle"] is not None:
            tblpr.append(keep["w:tblStyle"])
        b = '<w:{0} w:val="{1}" w:sz="{2}" w:space="0" w:color="000000"/>'
        borders = (b.format("top", "single", 12) + '<w:left w:val="nil"/>'
                   + b.format("bottom", "single", 12) + '<w:right w:val="nil"/>'
                   + b.format("insideH", "single", 4) + '<w:insideV w:val="nil"/>')
        for x in (f'<w:tblW {nsdecls("w")} w:w="5000" w:type="pct"/>',
                  f'<w:jc {nsdecls("w")} w:val="center"/>',
                  f'<w:tblBorders {nsdecls("w")}>{borders}</w:tblBorders>'):
            tblpr.append(parse_xml(x))
        rows = tbl.findall(qn("w:tr"))
        first_cell = rows[0].find(qn("w:tc")) if rows else None
        if first_cell is not None and "".join(first_cell.xpath(".//w:t/text()")).strip() == "Use case":
            # Bảng đặc tả use case: giữ đúng tỉ lệ cột của filters.lua, không để Word tự co cột nhãn.
            tblpr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
        for t in ("w:tblCellMar", "w:tblLook"):
            if keep[t] is not None:
                tblpr.append(keep[t])
        # Bảng ngắn: giữ trọn trên một trang (mẫu của Khoa, tr.19: không ngắt trang giữa bảng)
        if len(rows) <= 20:
            for r in rows[:-1]:
                for p in r.iter(qn("w:p")):
                    ppr = get_ppr(p)
                    if ppr.find(qn("w:keepNext")) is None:
                        insert_before_first(ppr, parse_xml(f"<w:keepNext {nsdecls('w')}/>"),
                                            ("keepLines", "pageBreakBefore", "framePr", "widowControl",
                                             "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs",
                                             "suppressAutoHyphens", "spacing", "ind", "jc", "outlineLvl", "rPr"))
        for r in rows:
            # không để một dòng bị cắt đôi giữa hai trang
            trpr = r.find(qn("w:trPr"))
            if trpr is None:
                trpr = parse_xml(f"<w:trPr {nsdecls('w')}/>")
                r.insert(1 if r.find(qn("w:tblPrEx")) is not None else 0, trpr)
            if trpr.find(qn("w:cantSplit")) is None:
                trpr.insert(0, parse_xml(f"<w:cantSplit {nsdecls('w')}/>"))
        header_rows = [r for r in rows if r.find(f"{qn('w:trPr')}/{qn('w:tblHeader')}") is not None]
        if not header_rows and rows:
            header_rows = rows[:1]
            trpr = rows[0].find(qn("w:trPr"))
            if trpr is None:
                trpr = parse_xml(f"<w:trPr {nsdecls('w')}/>")
                tcs = rows[0].find(qn("w:tblPrEx"))
                rows[0].insert(1 if tcs is not None else 0, trpr)
            trpr.append(parse_xml(f"<w:tblHeader {nsdecls('w')}/>"))
        for r in header_rows:
            for run in r.iter(qn("w:r")):
                rpr = run.find(qn("w:rPr"))
                if rpr is None:
                    rpr = parse_xml(f"<w:rPr {nsdecls('w')}/>")
                    run.insert(0, rpr)
                if rpr.find(qn("w:b")) is None:
                    insert_before_first(rpr, parse_xml(f"<w:b {nsdecls('w')}/>"),
                                        ("i", "iCs", "caps", "strike", "noProof", "vanish",
                                         "color", "spacing", "sz", "szCs", "highlight", "u",
                                         "shd", "vertAlign", "lang"))
    return n


EQ_TABLE = (
    f'<w:tbl {nsdecls("w")}><w:tblPr><w:tblW w:w="{{w}}" w:type="dxa"/><w:jc w:val="center"/>'
    '<w:tblBorders><w:top w:val="nil"/><w:left w:val="nil"/><w:bottom w:val="nil"/>'
    '<w:right w:val="nil"/><w:insideH w:val="nil"/><w:insideV w:val="nil"/></w:tblBorders>'
    '<w:tblLayout w:type="fixed"/><w:tblCellMar><w:left w:w="0" w:type="dxa"/>'
    '<w:right w:w="0" w:type="dxa"/></w:tblCellMar>'
    '<w:tblLook w:val="0000" w:firstRow="0" w:lastRow="0" w:firstColumn="0" w:lastColumn="0" w:noHBand="1" w:noVBand="1"/>'
    '</w:tblPr><w:tblGrid><w:gridCol w:w="{side}"/><w:gridCol w:w="{mid}"/><w:gridCol w:w="{side}"/></w:tblGrid>'
    '<w:tr><w:trPr><w:cantSplit/></w:trPr>'
    '<w:tc><w:tcPr><w:tcW w:w="{side}" w:type="dxa"/><w:vAlign w:val="center"/></w:tcPr>'
    '<w:p><w:pPr><w:pStyle w:val="Equation"/></w:pPr></w:p></w:tc>'
    '<w:tc><w:tcPr><w:tcW w:w="{mid}" w:type="dxa"/><w:vAlign w:val="center"/></w:tcPr></w:tc>'
    '<w:tc><w:tcPr><w:tcW w:w="{side}" w:type="dxa"/><w:vAlign w:val="center"/></w:tcPr></w:tc>'
    '</w:tr></w:tbl>')


def build_equation_tables(body) -> int:
    """Ghép đoạn công thức + đoạn số hiệu thành bảng 3 cột không viền: [ | công thức | (2-7)]."""
    total = round(15.5 / 2.54 * 1440)
    side = round(1.6 / 2.54 * 1440)
    xml = EQ_TABLE.replace("{w}", str(total)).replace("{side}", str(side)).replace("{mid}", str(total - 2 * side))
    n = 0
    for num_p in [p for p in body.iter(qn("w:p")) if pstyle(p) == "EquationNumber"]:
        math_p = num_p.getprevious()
        if math_p is None or math_p.find(f".//{{{NS_M}}}oMathPara") is None:
            warn(f"Không tìm thấy công thức đứng trước số hiệu {ptext(num_p)}")
            continue
        tbl = parse_xml(xml)
        math_p.addprevious(tbl)
        cells = tbl.findall(f".//{qn('w:tc')}")
        cells[1].append(math_p)
        cells[2].append(num_p)
        n += 1
    return n


def scale_images(body) -> None:
    """Thu ảnh về vừa khung: rộng tối đa TEXT_WIDTH, cao tối đa MAX_FIGURE_HEIGHT (kế hoạch mục 6.2)."""
    max_cx, max_cy = int(TEXT_WIDTH), int(MAX_FIGURE_HEIGHT)
    for inline in body.iter(f"{{{NS_WP}}}inline"):
        ext = inline.find(f"{{{NS_WP}}}extent")
        cx, cy = int(ext.get("cx")), int(ext.get("cy"))
        f = min(1.0, max_cx / cx, max_cy / cy)
        if f >= 1.0:
            continue
        new_cx, new_cy = round(cx * f), round(cy * f)
        ext.set("cx", str(new_cx))
        ext.set("cy", str(new_cy))
        for a_ext in inline.iter(f"{{{NS_A}}}ext"):
            if a_ext.get("cx"):
                a_ext.set("cx", str(new_cx))
                a_ext.set("cy", str(new_cy))


SECTPR_AFTER_PGNUM = ("cols", "formProt", "vAlign", "noEndnote", "titlePg", "textDirection",
                      "bidi", "rtlGutter", "docGrid", "printerSettings", "sectPrChange")


def page_field_runs() -> str:
    return ('<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            '<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>1</w:t></w:r>'
            '<w:r><w:fldChar w:fldCharType="end"/></w:r>')


def apply_sections(doc, names: list[str], cfg: dict) -> None:
    body = doc.element.body
    final = body.find(qn("w:sectPr"))
    for tag in ("w:headerReference", "w:footerReference", "w:pgNumType", "w:titlePg"):
        for el in final.findall(qn(tag)):
            final.remove(el)
    for p in [p for p in body.findall(qn("w:p")) if MARK_RE.fullmatch(ptext(p))]:
        for r in p.findall(qn("w:r")):
            p.remove(r)
        get_ppr(p).append(copy.deepcopy(final))

    sections = list(doc.sections)
    if len(sections) != len(names):
        warn(f"Số section ({len(sections)}) khác số phần ({len(names)})")
    numbering = {"cover": None, "front": "lowerRoman" if cfg.get("front_page_numbers") == "roman" else None,
                 "main": "decimal", "appendix": None}
    students = " – ".join(cfg.get("students", []))
    for sec, name in zip(sections, names):
        sec.start_type = WD_SECTION.NEW_PAGE
        fmt = numbering.get(name)
        sp = sec._sectPr
        for el in sp.findall(qn("w:pgNumType")):
            sp.remove(el)
        if fmt:
            insert_before_first(sp, parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="{fmt}" w:start="1"/>'),
                                SECTPR_AFTER_PGNUM)

        sec.header.is_linked_to_previous = False
        hp = sec.header.paragraphs[0]._p
        for r in hp.findall(qn("w:r")):
            hp.remove(r)
        set_pstyle(hp, "Header")
        get_ppr(hp).append(parse_xml(f'<w:jc {nsdecls("w")} w:val="{cfg.get("header_align", "right")}"/>'))
        italic = "<w:rPr><w:i/></w:rPr>" if cfg.get("header_italic", True) else ""
        if name != "cover":   # bìa và phụ bìa không có đầu trang, chân trang
            hp.append(parse_xml(f'<w:r {nsdecls("w")}>{italic}<w:t xml:space="preserve">{cfg["header_text"]}</w:t></w:r>'))

        sec.footer.is_linked_to_previous = False
        fp = sec.footer.paragraphs[0]._p
        for r in fp.findall(qn("w:r")):
            fp.remove(r)
        set_pstyle(fp, "Footer")
        runs = "" if name == "cover" else f'<w:r><w:t xml:space="preserve">{students}</w:t></w:r>'
        if fmt:
            runs += "<w:r><w:tab/></w:r>" + page_field_runs()
        for r in parse_xml(f'<w:p {nsdecls("w")}>{runs}</w:p>'):
            fp.append(r)


SETTINGS_AFTER_UPDATE = ("hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids",
                         "mathPr", "attachedSchema", "themeFontLang", "clrSchemeMapping",
                         "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
                         "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary",
                         "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator")


def set_update_fields(doc, on: bool) -> None:
    settings = doc.settings.element
    for el in settings.findall(qn("w:updateFields")):
        settings.remove(el)
    if on:
        insert_before_first(settings, parse_xml(f'<w:updateFields {nsdecls("w")} w:val="true"/>'),
                            SETTINGS_AFTER_UPDATE)


# ---------------------------------------------------------------------------
# 3. Kiểm tra
# ---------------------------------------------------------------------------
def check_mentions(body) -> dict:
    """Mọi Hình/Bảng phải được nhắc trong đoạn văn trước khi xuất hiện (mẫu của Khoa)."""
    seen_text: list[str] = []
    stats = {"hinh": 0, "bang": 0, "todo": 0}
    for p in body.iter(qn("w:p")):
        st, text = pstyle(p), ptext(p)
        if st == "Caption":
            m = re.match(r"(Hình|Bảng) (\d+\.\d+)", text)
            if m:
                stats["hinh" if m.group(1) == "Hình" else "bang"] += 1
                ref = f"{m.group(1)} {m.group(2)}"
                if not any(re.search(re.escape(ref) + r"(?!\d)", t) for t in seen_text):
                    warn(f"{ref} chưa được nhắc trong văn bản trước khi xuất hiện")
        elif st == "TODO":
            stats["todo"] += 1
        else:
            seen_text.append(text)
    return stats


def report_page_budget(pdf: Path, budget: dict) -> None:
    """In số trang thực tế của từng chương trong PDF và so với ngân sách (KE_HOACH_V2.md mục 2)."""
    try:
        import pymupdf
    except ImportError:
        print("  [ngân sách trang] cần gói pymupdf để đếm trang")
        return
    starts: dict[str, int] = {}
    with pymupdf.open(str(pdf)) as d:
        n_pages = d.page_count
        for i, page in enumerate(d):
            for line in page.get_text().splitlines():
                line = line.strip()
                # Dòng tiêu đề thật, không phải dòng mục lục (dòng mục lục kết thúc bằng số trang)
                m = re.match(r"^CHƯƠNG (\d+)\b", line)
                if m and not re.search(r"\d\s*$", line[len(m.group(0)):]):
                    starts[m.group(1)] = i
                elif line == "TÀI LIỆU THAM KHẢO":
                    starts["TLTK"] = i
                elif line == "PHỤ LỤC":
                    starts["PL"] = i
    order = sorted((k for k in starts if k.isdigit()), key=int) + [k for k in ("TLTK", "PL") if k in starts]
    if not order:
        return
    print("  Ngân sách trang (phần chính):")
    for idx, k in enumerate(order):
        if k == "PL":
            break
        end = starts[order[idx + 1]] if idx + 1 < len(order) else n_pages
        pages, cap = end - starts[k], budget.get(k)
        name = "Tài liệu tham khảo" if k == "TLTK" else f"Chương {k}"
        flag = "" if cap is None else ("  VƯỢT" if pages > cap * 1.1 else "")
        warn_txt = f" / {cap}" if cap is not None else ""
        print(f"    {name:20} {pages:4}{warn_txt}{flag}")
    main_end = starts.get("PL", n_pages)
    main_pages = main_end - starts[order[0]]
    cap = budget.get("main")
    print(f"    {'Phần chính':20} {main_pages:4}" + (f" / {cap}" + ("  VƯỢT" if main_pages > cap else "") if cap else ""))


# ---------------------------------------------------------------------------
def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--sample", action="store_true", help="ghép file mẫu kiểm tra định dạng (build/sample)")
    ap.add_argument("--word", action="store_true", help="dùng Word cập nhật mục lục, caption, số trang")
    ap.add_argument("--pdf", action="store_true", help="xuất thêm PDF (cần --word)")
    ap.add_argument("--config", default=str(BUILD / "report_config.json"))
    args = ap.parse_args()

    cfg = json.loads(Path(args.config).read_text(encoding="utf-8"))
    parts = cfg["sample_parts"] if args.sample else cfg["parts"]
    out = REPORT / (cfg["sample_output"] if args.sample else cfg["output"])

    markdown, names = assemble(parts)
    markdown = fill_placeholders(markdown, cfg)
    if not names:
        print("Không có bản thảo nào để ghép.")
        return 1
    raw_docx = BUILD / "_pandoc_raw.docx"
    run_pandoc(markdown, raw_docx, cfg)

    doc = docx.Document(str(raw_docx))
    body = doc.element.body
    add_heading_numbering(doc, cfg.get("heading1_layout", "one_line_left"))
    restyle_paragraphs(body)
    n_tbl = format_tables(body)
    n_eq = build_equation_tables(body)
    scale_images(body)
    apply_sections(doc, names, cfg)
    set_update_fields(doc, True)
    stats = check_mentions(body)
    doc.save(str(out))
    raw_docx.unlink(missing_ok=True)

    if args.word:
        pdf = out.with_suffix(".pdf") if args.pdf else None
        cmd = ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File",
               str(BUILD / "word_update.ps1"), "-Docx", str(out)]
        if pdf:
            cmd += ["-Pdf", str(pdf)]
        subprocess.run(cmd, check=True)
        # Trường đã được Word cập nhật: tắt hộp thoại "update fields" khi mở file
        d2 = docx.Document(str(out))
        set_update_fields(d2, False)
        d2.save(str(out))
        if pdf:
            print(f"PDF: {pdf}")
            report_page_budget(pdf, cfg.get("page_budget", {}))

    print(f"Đã ghép: {out}")
    print(f"  phần: {', '.join(names)} | hình: {stats['hinh']} | bảng: {stats['bang']} "
          f"(tổng {n_tbl} bảng) | công thức đánh số: {n_eq} | ghi chú TODO: {stats['todo']}")
    if warnings:
        print(f"  {len(warnings)} cảnh báo (xem ở trên)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
